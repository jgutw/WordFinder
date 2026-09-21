#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Oxford 3000 full run: Version B only.
Processes all 3000 words across all languages. Resumable via checkpoint.
Does not modify the dashboard or web app.
"""

import os
import re
import time
import json
from pathlib import Path

import sys
sys.path.insert(0, str(Path(__file__).resolve().parent))

from wordfinder import (
    LANGS, normalize_text, Translator, postprocess_translation_for_lang,
    analyze_translation, PER_LANG_ZIPF, DEFAULT_ZIPF, DEFAULT_MIN_LEN_LATINLIKE
)
from dataclasses import dataclass

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Install openpyxl: pip install openpyxl")
    sys.exit(1)


CHECKPOINT_INTERVAL = 10  # Save progress every N words
CHECKPOINT_FILE = Path(__file__).parent / "data" / "oxford_3000_checkpoint.json"
OUTPUT_FILE = Path(__file__).parent / "data" / "oxford_3000_version_b.xlsx"


@dataclass
class AnalysisArgs:
    zh_tokenizer: str = "jieba"
    jp_tokenizer: str = "tiny"
    th_tokenizer: str = "pythainlp"


def extract_words_from_text(text: str) -> list[str]:
    """Extract headwords from Oxford 3000 raw text."""
    words = []
    for line in text.splitlines():
        line = line.strip()
        if not line or not re.search(r'\s+[AB][12]\s*$', line):
            continue
        head = re.sub(r'\s+[AB][12]\s*$', '', line)
        head = re.sub(r'\s+(?:n\.|v\.|adj\.|adv\.|prep\.|conj\.|det\.|pron\.|number|'
                      r'indefinite article|definite article|modal v\.|auxiliary v\.|'
                      r'exclam\.|adj\./adv\.|adv\./n\.|infinitive marker).*$', '', head)
        head = re.sub(r'\s*\([^)]*\)', '', head).strip().split(',')[0].strip()
        if head and len(head) > 1:
            words.append(head)
    seen = set()
    return [w for w in words if w.lower() not in seen and not seen.add(w.lower())]


def has_inner_words(result: dict) -> bool:
    main = result.get('inner_words_main', [])
    rare = result.get('inner_words_rare', [])
    return len(main) > 0 or len(rare) > 0


def format_inner_words(main: list, rare: list) -> str:
    parts = []
    for it in main + rare:
        gloss = it.get('gloss_en', '') or ''
        parts.append(f"{it['token']} → {gloss}")
    return "; ".join(parts)


def process_word(english: str, translator, args, langs) -> dict:
    """Process one word across all languages. Returns {lang_name: {...}} only for langs with inner words."""
    results = {}
    per_lang_zipf = dict(PER_LANG_ZIPF)
    for i, lang in enumerate(langs):
        try:
            raw = translator.to_target(english, target=lang.trans_code)
            translated = postprocess_translation_for_lang(raw, lang, do_lemma=False)
            result = analyze_translation(
                translated, lang, DEFAULT_ZIPF, per_lang_zipf,
                DEFAULT_MIN_LEN_LATINLIKE, 12, translator, True, args=args
            )
            if has_inner_words(result):
                results[lang.name] = {
                    'translation': result['translation'],
                    'inner_words_main': result['inner_words_main'],
                    'inner_words_rare': result.get('inner_words_rare', []),
                }
        except Exception as e:
            print(f"  Error {lang.name}: {e}", flush=True)
        if (i + 1) % 10 == 0:
            time.sleep(0.5)
        time.sleep(0.3)
    return results


def save_checkpoint(word_data: dict, next_index: int, words: list):
    """Save progress to JSON for resumption (atomic write to avoid partial files on crash)."""
    Path(CHECKPOINT_FILE).parent.mkdir(parents=True, exist_ok=True)
    data = {
        'word_data': word_data,
        'next_index': next_index,
        'total_words': len(words),
    }
    tmp = CHECKPOINT_FILE.with_suffix('.json.tmp')
    with open(tmp, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=0)
        f.flush()
        os.fsync(f.fileno())
    os.replace(tmp, CHECKPOINT_FILE)
    print(f"  Checkpoint saved ({next_index}/{len(words)} done)", flush=True)


def load_checkpoint() -> tuple[dict, int, list] | None:
    """Load checkpoint if exists. Returns (word_data, next_index, words) or None."""
    if not CHECKPOINT_FILE.exists():
        return None
    try:
        raw = CHECKPOINT_FILE.read_text(encoding='utf-8')
        try:
            data = json.loads(raw)
        except json.JSONDecodeError:
            dec = json.JSONDecoder()
            data, end = dec.raw_decode(raw)
            tail = raw[end:].strip()
            if tail:
                print(
                    f"Checkpoint had {len(tail)} chars of trailing data (e.g. interrupted save); trimmed.",
                    flush=True,
                )
                CHECKPOINT_FILE.write_text(raw[:end], encoding='utf-8')
        raw_path = Path(__file__).parent / "oxford_3000_raw.txt"
        text = raw_path.read_text(encoding='utf-8') if raw_path.exists() else ""
        words = extract_words_from_text(text)
        return data['word_data'], data['next_index'], words
    except Exception as e:
        print(f"Checkpoint load failed: {e}", flush=True)
        return None


def write_version_b(word_data: dict, out_path: Path):
    """Version B: One row per word, one column per language."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Oxford 3000"
    all_langs = sorted({lang for data in word_data.values() for lang in data})
    ws.cell(1, 1, "English Word").font = Font(bold=True)
    for c, lang in enumerate(all_langs, 2):
        ws.cell(1, c, lang).font = Font(bold=True)
    for r, (english, data) in enumerate(word_data.items(), 2):
        ws.cell(r, 1, english)
        for c, lang in enumerate(all_langs, 2):
            if lang in data:
                d = data[lang]
                cell_val = f"{d['translation']} | Inner: {format_inner_words(d['inner_words_main'], d['inner_words_rare'])}"
                ws.cell(r, c, cell_val)
    for col in range(1, len(all_langs) + 2):
        ws.column_dimensions[get_column_letter(col)].width = 35
    wb.save(out_path)
    print(f"Saved Version B: {out_path}", flush=True)


def main():
    raw_path = Path(__file__).parent / "oxford_3000_raw.txt"
    if not raw_path.exists():
        print("Missing oxford_3000_raw.txt", flush=True)
        sys.exit(1)
    words = extract_words_from_text(raw_path.read_text(encoding='utf-8'))
    total = len(words)
    print(f"Oxford 3000 full run (Version B): {total} words, {len(LANGS)} languages", flush=True)

    # Try to resume from checkpoint
    checkpoint = load_checkpoint()
    if checkpoint:
        word_data, start_index, ck_words = checkpoint
        if ck_words == words and start_index < total:
            print(f"Resuming from word {start_index + 1}/{total}", flush=True)
        else:
            word_data = {}
            start_index = 0
    else:
        word_data = {}
        start_index = 0

    translator = Translator(default_source="en", delay=0.4)
    args = AnalysisArgs()
    langs = LANGS

    for i in range(start_index, total):
        english = words[i]
        print(f"[{i+1}/{total}] {english}...", flush=True)
        data = process_word(english, translator, args, langs)
        word_data[english] = data

        if (i + 1) % CHECKPOINT_INTERVAL == 0:
            save_checkpoint(word_data, i + 1, words)

    Path(OUTPUT_FILE).parent.mkdir(parents=True, exist_ok=True)
    write_version_b(word_data, OUTPUT_FILE)
    if CHECKPOINT_FILE.exists():
        CHECKPOINT_FILE.unlink()
        print("Checkpoint removed.", flush=True)
    print("Done.", flush=True)


if __name__ == "__main__":
    main()
