#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Oxford 3000 sample: Process 10 random words, output Version A and B Excel files.
"""

import re
import random
import time
import json
from pathlib import Path

# Add project root for imports
import sys
sys.path.insert(0, str(Path(__file__).resolve().parent))

from wordfinder import (
    LANGS, normalize_text, Translator, postprocess_translation_for_lang,
    analyze_translation, PER_LANG_ZIPF, DEFAULT_ZIPF, DEFAULT_MIN_LEN_LATINLIKE
)
from dataclasses import dataclass

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Install openpyxl: pip install openpyxl")
    sys.exit(1)


@dataclass
class AnalysisArgs:
    zh_tokenizer: str = "jieba"
    jp_tokenizer: str = "tiny"
    th_tokenizer: str = "pythainlp"


def extract_words_from_text(text: str) -> list[str]:
    """Extract headwords from Oxford 3000 raw text."""
    words = []
    for line in text.splitlines():
        line = line.strip()
        if not line or not re.search(r'\s+[AB][12]\s*$', line):
            continue
        head = re.sub(r'\s+[AB][12]\s*$', '', line)
        head = re.sub(r'\s+(?:n\.|v\.|adj\.|adv\.|prep\.|conj\.|det\.|pron\.|number|'
                      r'indefinite article|definite article|modal v\.|auxiliary v\.|'
                      r'exclam\.|adj\./adv\.|adv\./n\.|infinitive marker).*$', '', head)
        head = re.sub(r'\s*\([^)]*\)', '', head).strip().split(',')[0].strip()
        if head and len(head) > 1:
            words.append(head)
    seen = set()
    return [w for w in words if w.lower() not in seen and not seen.add(w.lower())]


def has_inner_words(result: dict) -> bool:
    main = result.get('inner_words_main', [])
    rare = result.get('inner_words_rare', [])
    return len(main) > 0 or len(rare) > 0


def format_inner_words(main: list, rare: list) -> str:
    parts = []
    for it in main + rare:
        gloss = it.get('gloss_en', '') or ''
        parts.append(f"{it['token']} → {gloss}")
    return "; ".join(parts)


def process_word(english: str, translator, args, langs) -> dict:
    """Process one word across all languages. Returns {lang_name: {...}} only for langs with inner words."""
    results = {}
    per_lang_zipf = dict(PER_LANG_ZIPF)
    for i, lang in enumerate(langs):
        try:
            raw = translator.to_target(english, target=lang.trans_code)
            translated = postprocess_translation_for_lang(raw, lang, do_lemma=False)
            result = analyze_translation(
                translated, lang, DEFAULT_ZIPF, per_lang_zipf,
                DEFAULT_MIN_LEN_LATINLIKE, 12, translator, True, args=args
            )
            if has_inner_words(result):
                results[lang.name] = {
                    'translation': result['translation'],
                    'inner_words_main': result['inner_words_main'],
                    'inner_words_rare': result.get('inner_words_rare', []),
                }
        except Exception as e:
            print(f"  Error {lang.name}: {e}")
        if (i + 1) % 10 == 0:
            time.sleep(0.5)
        time.sleep(0.3)
    return results


def write_version_a(rows: list, out_path: Path):
    """Version A: One row per (English word, language) with inner words."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Version A"
    headers = ["English Word", "Language", "Translation", "Inner Words (token → English)"]
    for c, h in enumerate(headers, 1):
        ws.cell(1, c, h).font = Font(bold=True)
    for r, row in enumerate(rows, 2):
        ws.cell(r, 1, row['english'])
        ws.cell(r, 2, row['language'])
        ws.cell(r, 3, row['translation'])
        ws.cell(r, 4, row['inner_words_str'])
    for col in range(1, 5):
        ws.column_dimensions[get_column_letter(col)].width = 25
    wb.save(out_path)
    print(f"Saved Version A: {out_path}")


def write_version_b(word_data: dict, out_path: Path):
    """Version B: One row per word, one column per language."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Version B"
    all_langs = sorted({lang for data in word_data.values() for lang in data})
    ws.cell(1, 1, "English Word").font = Font(bold=True)
    for c, lang in enumerate(all_langs, 2):
        ws.cell(1, c, lang).font = Font(bold=True)
    for r, (english, data) in enumerate(word_data.items(), 2):
        ws.cell(r, 1, english)
        for c, lang in enumerate(all_langs, 2):
            if lang in data:
                d = data[lang]
                cell_val = f"{d['translation']} | Inner: {format_inner_words(d['inner_words_main'], d['inner_words_rare'])}"
                ws.cell(r, c, cell_val)
    for col in range(1, len(all_langs) + 2):
        ws.column_dimensions[get_column_letter(col)].width = 35
    wb.save(out_path)
    print(f"Saved Version B: {out_path}")


def main():
    raw_path = Path(__file__).parent / "oxford_3000_raw.txt"
    if raw_path.exists():
        text = raw_path.read_text(encoding='utf-8')
    else:
        text = "ability n. A2\nbeautiful adj. A1\nchocolate n. A1\ncomputer n. A1\ndictionary n. A1\nfriend n. A1\nhospital n. A1\nlanguage n. A1\nmountain n. A1\nrestaurant n. A1\n"
    words = extract_words_from_text(text)
    if len(words) < 10:
        words = ["ability", "beautiful", "chocolate", "computer", "dictionary",
                 "friend", "hospital", "language", "mountain", "restaurant"]
    random.seed(42)
    sample = random.sample(words, min(10, len(words)))
    print(f"Processing 10 words: {sample}", flush=True)

    translator = Translator(default_source="en", delay=0.4)
    args = AnalysisArgs()
    langs = LANGS  # all languages

    word_data = {}
    version_a_rows = []

    for i, english in enumerate(sample):
        print(f"[{i+1}/10] {english}...", flush=True)
        data = process_word(english, translator, args, langs)
        word_data[english] = data
        for lang_name, d in data.items():
            version_a_rows.append({
                'english': english,
                'language': lang_name,
                'translation': d['translation'],
                'inner_words_str': format_inner_words(d['inner_words_main'], d['inner_words_rare']),
            })

    out_dir = Path(__file__).parent / "data"
    out_dir.mkdir(exist_ok=True)
    write_version_a(version_a_rows, out_dir / "oxford_3000_sample_version_a.xlsx")
    write_version_b(word_data, out_dir / "oxford_3000_sample_version_b.xlsx")
    print("Done.")


if __name__ == "__main__":
    main()
